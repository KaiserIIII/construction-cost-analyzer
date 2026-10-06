#!/usr/bin/env python3
"""Import a reviewed ONS OPI release; refuse a changed upstream workbook.

Reading the checked-in CSV needs only Python's standard library. Rebuilding it
requires the already installed, optional openpyxl XLSX reader.
"""

import argparse
import csv
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
import hashlib
import io
import json
from pathlib import Path
import platform
import re
import urllib.error
import urllib.request


SOURCE_URL = "https://www.ons.gov.uk/businessindustryandtrade/constructionindustry/datasets/interimconstructionoutputpriceindices"
WORKBOOK_URL = "https://www.ons.gov.uk/file?uri=%2Fbusinessindustryandtrade%2Fconstructionindustry%2Fdatasets%2Finterimconstructionoutputpriceindices%2Fcurrent%2Fbulletindataset9.xlsx"
PINNED_SHA256 = "ea0cbfe6573be52210ea0469f182ac5f03c68af39b193ab0f328feb24626edde"
RELEASE_DATE = "2026-08-13"
SOURCE_REF = "ons_construction_opi_2026-08-13"
CSV_FILENAME = "ons_construction_opi.csv"
CSV_FIELDS = ("period", "index_type", "series", "index_value", "base_year", "source_type", "source_ref")
SHEETS = ("All construction", "New work", "Repair & maintenance")
MONTHS = {name: number for number, name in enumerate(
    ("jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"), 1)}


def _text(value):
    return " ".join(str(value or "").split())


def _period(value, year):
    """Recognise monthly rows, carrying the year between January labels."""
    if isinstance(value, (date, datetime)):
        return f"{value.year:04d}-{value.month:02d}", value.year
    label = _text(value)
    if not label:
        return None, year
    explicit_year = re.match(r"^(\d{4})(?:\s+|$)", label)
    if explicit_year:
        year = int(explicit_year.group(1))
        label = label[explicit_year.end():].strip()
    if not label or re.fullmatch(r"(?:Q[1-4]|Quarter\s*[1-4]|Annual|Year|Total)", label, flags=re.I):
        return None, year
    month = MONTHS.get(label[:3].lower())
    if month is None or not re.fullmatch(r"[A-Za-z]+", label):
        raise ValueError(f"Unrecognised time period: {value!r}")
    full_month_names = ("january", "february", "march", "april", "may", "june", "july",
                        "august", "september", "october", "november", "december")
    if label.lower() not in MONTHS and label.lower() not in full_month_names:
        raise ValueError(f"Unrecognised time period: {value!r}")
    if year is None:
        raise ValueError(f"Monthly label has no year: {value!r}")
    return f"{year:04d}-{month:02d}", year


def _level(value, sheet, row_number, series):
    try:
        if isinstance(value, bool) or value is None:
            raise InvalidOperation
        level = Decimal(str(value))
        if not level.is_finite() or level <= 0:
            raise InvalidOperation
    except (InvalidOperation, ValueError):
        raise ValueError(f"Invalid index level at {sheet} row {row_number}, {series}: {value!r}") from None
    result = format(level, "f")
    return result if "." in result else result + ".0"


def parse_workbook(content, *, expected_sha256=PINNED_SHA256):
    """Return sorted unique monthly index levels, after checking source bytes.

    expected_sha256 exists for controlled layout tests. The command-line import
    always uses the reviewed release hash and has no digest override.
    """
    actual_sha256 = hashlib.sha256(content).hexdigest()
    if actual_sha256 != expected_sha256:
        raise ValueError(f"Workbook SHA-256 mismatch: expected {expected_sha256}, received {actual_sha256}. "
                         "The current ONS release may have changed; review it before updating the pin.")
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise RuntimeError("Rebuilding the ONS snapshot requires optional openpyxl; the CSV needs no XLSX dependency.") from None

    book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    records = {}
    origins = {}
    try:
        found_sheet = False
        for sheet in book:
            if sheet.title not in SHEETS:
                continue
            found_sheet = True
            header_found = False
            columns = []
            period_column = None
            year_column = None
            year = None
            for row_number, cells in enumerate(sheet.iter_rows(values_only=True), 1):
                if not header_found:
                    labels = [_text(cell) for cell in cells]
                    candidates = [i for i, label in enumerate(labels) if label.lower() in ("time period", "period", "month")]
                    if not candidates:
                        continue
                    period_column = candidates[0]
                    year_columns = [i for i, label in enumerate(labels) if label.lower() == "year"]
                    year_column = year_columns[0] if year_columns else None
                    for column, label in enumerate(labels):
                        index_word = re.search(r"\bindex\b", label, flags=re.I)
                        base = re.search(r"(\d{4})\s*=\s*100", label)
                        if index_word and base and not re.search(r"percentage|%", label, flags=re.I):
                            series = label[:index_word.start()].strip(" ,(")
                            if not series:
                                raise ValueError(f"Missing series label in {sheet.title}")
                            columns.append((column, series, base.group(1)))
                    if not columns:
                        raise ValueError(f"No index-level columns in {sheet.title}")
                    header_found = True
                    continue
                if not any(cell is not None for cell in cells):
                    continue
                if year_column is not None and cells[year_column] is not None:
                    separate_year = _text(cells[year_column])
                    if not re.fullmatch(r"\d{4}", separate_year):
                        raise ValueError(f"Invalid year in {sheet.title} row {row_number}: {separate_year!r}")
                    year = int(separate_year)
                period, year = _period(cells[period_column], year)
                if period is None:
                    continue
                for column, series, base_year in columns:
                    record = {
                        "period": period, "index_type": "OPI", "series": series,
                        "index_value": _level(cells[column], sheet.title, row_number, series),
                        "base_year": base_year, "source_type": "official_statistics", "source_ref": SOURCE_REF,
                    }
                    key = (period, series)
                    if key in records:
                        if origins[key] == sheet.title:
                            raise ValueError(f"Duplicate monthly observation in {sheet.title}: {key}")
                        if record != records[key]:
                            raise ValueError(f"Conflicting duplicate index level across sheets: {key}")
                    else:
                        records[key] = record
                        origins[key] = sheet.title
            if not header_found:
                raise ValueError(f"No time-period header in {sheet.title}")
        if not found_sheet or not records:
            raise ValueError("No monthly ONS index levels found")
    finally:
        book.close()
    return [records[key] for key in sorted(records)]


def import_snapshot(content, output_dir):
    """Verify the pinned release and write deterministic CSV plus provenance."""
    rows = parse_workbook(content)
    months = sorted({row["period"] for row in rows})
    series = sorted({row["series"] for row in rows})
    expected_months = [f"{year:04d}-{month:02d}" for year in range(2014, 2027)
                       for month in range(1, 13) if (year, month) <= (2026, 6)]
    if months != expected_months or len(series) != 10 or len(rows) != 1500:
        raise ValueError("Pinned release must contain 150 months, 10 series and 1,500 unique index levels")
    for name in series:
        if [row["period"] for row in rows if row["series"] == name] != expected_months:
            raise ValueError(f"Incomplete monthly coverage for {name}")
    if {row["base_year"] for row in rows} != {"2015"}:
        raise ValueError("Pinned release must use 2015=100")

    stream = io.StringIO(newline="")
    writer = csv.DictWriter(stream, fieldnames=CSV_FIELDS, lineterminator="\n")
    writer.writeheader()
    writer.writerows(rows)
    csv_bytes = stream.getvalue().encode("utf-8")
    import openpyxl
    manifest = {
        "dataset_title": "ONS Construction Output Price Indices (OPIs)",
        "source_ref": SOURCE_REF,
        "source_type": "official_statistics",
        "source_url": SOURCE_URL,
        "workbook_url": WORKBOOK_URL,
        "release_date": RELEASE_DATE,
        "retrieved_at_utc": datetime.now(timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z"),
        "licence": "Open Government Licence v3.0",
        "licence_url": "https://www.nationalarchives.gov.uk/doc/open-government-licence/version/3/",
        "attribution": "Contains public sector information licensed under the Open Government Licence v3.0. Source: Office for National Statistics. Crown copyright 2026.",
        "coverage_workbook": "Great Britain",
        "coverage_dataset_page": "UK",
        "coverage_note": "The source workbook cover says Great Britain; the ONS dataset page describes UK. Both source labels are retained without harmonising them.",
        "index_type": "OPI",
        "base_year": 2015,
        "base_description": "2015=100",
        "seasonal_adjustment": "not seasonally adjusted",
        "observation_frequency": "monthly",
        "release_frequency": "quarterly",
        "period_start": months[0],
        "period_end": months[-1],
        "row_count": len(rows),
        "month_count": len(months),
        "series_count": len(series),
        "series": series,
        "source_worksheets": list(SHEETS),
        "duplicate_aggregate_observations_removed": 300,
        "exclusions": ["Percentage-change columns", "Annual and quarterly period rows", "Repeated cross-sheet aggregate levels after equality checks"],
        "workbook_sha256": hashlib.sha256(content).hexdigest(),
        "workbook_bytes": len(content),
        "csv_file": CSV_FILENAME,
        "csv_sha256": hashlib.sha256(csv_bytes).hexdigest(),
        "csv_encoding": "UTF-8",
        "csv_line_endings": "LF",
        "csv_fields": list(CSV_FIELDS),
        "importer": "scripts/import_ons_opi.py",
        "import_environment": {"python": platform.python_version(), "openpyxl": openpyxl.__version__},
        "limitations": ["Output price index, not BCIS Tender Price Index (TPI)", "Historical index levels, not inflation forecasts", "Not individual project observations or a market quote", "No regional location factors, currency conversion or project-specific calibration", "ONS may revise historical data; this is a frozen reviewed release"],
    }
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    (output_dir / CSV_FILENAME).write_bytes(csv_bytes)
    (output_dir / "ons_construction_opi.metadata.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")
    return manifest


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, help="Use a local copy of the pinned public workbook")
    parser.add_argument("--output-dir", type=Path,
                        default=Path(__file__).resolve().parents[1] / "data" / "indices")
    args = parser.parse_args()
    try:
        if args.workbook:
            content = args.workbook.read_bytes()
        else:
            request = urllib.request.Request(WORKBOOK_URL, headers={"User-Agent": "Mozilla/5.0"})
            with urllib.request.urlopen(request, timeout=45) as response:
                content = response.read()
        manifest = import_snapshot(content, args.output_dir)
    except (ValueError, RuntimeError, OSError, urllib.error.URLError) as error:
        parser.error(str(error))
    print(f"Imported {manifest['row_count']} OPI levels: {manifest['month_count']} months, {manifest['series_count']} series.")
    print(f"CSV SHA-256: {manifest['csv_sha256']}")


if __name__ == "__main__":
    main()
